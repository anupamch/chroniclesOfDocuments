"""
Production-grade document parser with multiple OCR backends
"""
import os
from pathlib import Path
from typing import Literal
import pypdf
import pdfplumber
import docx
from app.core.config import settings


class ProductionDocumentParser:
    """Production-ready parser with fallback strategies"""
    
    OCR_BACKEND: Literal["tesseract", "cloud", "vision"] = settings.ocr_method
    
    @staticmethod
    def parse_pdf(file_path: str) -> tuple[str, list[dict]]:
        """Extract text and tables from PDF with fallback"""
        text = ""
        tables = []
        
        # Try direct text extraction first
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                text += page_text + "\n"
        
        # If scanned PDF, use appropriate OCR backend
        if len(text.strip()) < 50:
            print("  [Parser] Detected scanned PDF, using OCR...")
            
            if ProductionDocumentParser.OCR_BACKEND == "tesseract":
                text = ProductionDocumentParser._ocr_tesseract_pdf(file_path)
            elif ProductionDocumentParser.OCR_BACKEND == "cloud":
                text = ProductionDocumentParser._ocr_cloud_pdf(file_path)
            else:  # vision fallback
                text = ProductionDocumentParser._ocr_vision_pdf(file_path)
        
        # Extract tables
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_tables = page.extract_tables()
                    if page_tables:
                        tables.extend([{"data": table} for table in page_tables])
        except:
            pass
        
        return text.strip(), tables
    
    @staticmethod
    def parse_docx(file_path: str) -> tuple[str, list[dict]]:
        """Extract text and tables from DOCX"""
        doc = docx.Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
        
        tables = []
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                table_data.append([cell.text for cell in row.cells])
            tables.append({"data": table_data})
        
        return text.strip(), tables
    
    @staticmethod
    def parse_image(file_path: str) -> tuple[str, list[dict]]:
        """Extract text from image using configured backend"""
        
        if ProductionDocumentParser.OCR_BACKEND == "tesseract":
            text = ProductionDocumentParser._ocr_tesseract_image(file_path)
        elif ProductionDocumentParser.OCR_BACKEND == "cloud":
            text = ProductionDocumentParser._ocr_cloud_image(file_path)
        else:
            text = ProductionDocumentParser._ocr_vision_image(file_path)
        
        return text.strip(), []
    
    # ===== TESSERACT OCR (Best for production) =====
    
    @staticmethod
    def _ocr_tesseract_pdf(file_path: str) -> str:
        """OCR PDF using Tesseract (fast, accurate, scalable)"""
        import pdf2image
        import pytesseract
        
        images = pdf2image.convert_from_path(file_path)
        text = ""
        
        for i, image in enumerate(images):
            print(f"  [Tesseract] Page {i+1}/{len(images)}")
            custom_config = r'--oem 3 --psm 6'
            page_text = pytesseract.image_to_string(image, lang='eng+ben', config=custom_config)
            text += page_text + "\n"
        
        return text
    
    @staticmethod
    def _ocr_tesseract_image(file_path: str) -> str:
        """OCR image using Tesseract"""
        import pytesseract
        from PIL import Image
        
        image = Image.open(file_path)
        if image.mode != 'RGB':
            image = image.convert('RGB')
        
        custom_config = r'--oem 3 --psm 6'
        return pytesseract.image_to_string(image, lang='eng+ben', config=custom_config)
    
    # ===== CLOUD OCR (Best accuracy, costs money) =====
    
    @staticmethod
    def _ocr_cloud_pdf(file_path: str) -> str:
        """OCR using cloud services (AWS Textract, Google Vision, Azure)"""
        # Example: AWS Textract
        try:
            import boto3
            
            textract = boto3.client('textract')
            
            with open(file_path, 'rb') as document:
                response = textract.detect_document_text(
                    Document={'Bytes': document.read()}
                )
            
            text = ""
            for item in response['Blocks']:
                if item['BlockType'] == 'LINE':
                    text += item['Text'] + "\n"
            
            return text
            
        except Exception as e:
            print(f"  [Cloud OCR] Error: {e}")
            # Fallback to tesseract
            return ProductionDocumentParser._ocr_tesseract_pdf(file_path)
    
    @staticmethod
    def _ocr_cloud_image(file_path: str) -> str:
        """OCR image using cloud services"""
        # Similar to _ocr_cloud_pdf but for images
        try:
            import boto3
            
            textract = boto3.client('textract')
            
            with open(file_path, 'rb') as document:
                response = textract.detect_document_text(
                    Document={'Bytes': document.read()}
                )
            
            text = ""
            for item in response['Blocks']:
                if item['BlockType'] == 'LINE':
                    text += item['Text'] + "\n"
            
            return text
            
        except Exception as e:
            print(f"  [Cloud OCR] Error: {e}")
            return ProductionDocumentParser._ocr_tesseract_image(file_path)
    
    # ===== VISION MODEL OCR (Development/fallback) =====
    
    @staticmethod
    def _ocr_vision_pdf(file_path: str) -> str:
        """OCR using Ollama vision (fallback to Tesseract on failure)"""
        try:
            from app.core.parsers_vision import VisionDocumentParser
            text, _ = VisionDocumentParser.parse_pdf(file_path)
            if len(text.strip()) < 5:
                raise ValueError("Vision OCR returned empty text")
            return text
        except Exception as e:
            print(f"  [Vision OCR Fallback] Error: {e}")
            return ProductionDocumentParser._ocr_tesseract_pdf(file_path)
    
    @staticmethod
    def _ocr_vision_image(file_path: str) -> str:
        """OCR using Ollama vision (fallback to Tesseract on failure)"""
        try:
            from app.core.parsers_vision import VisionDocumentParser
            text, _ = VisionDocumentParser.parse_image(file_path)
            if len(text.strip()) < 5:
                raise ValueError("Vision OCR returned empty text")
            return text
        except Exception as e:
            print(f"  [Vision OCR Fallback] Error: {e}")
            return ProductionDocumentParser._ocr_tesseract_image(file_path)
    
    @staticmethod
    def parse_document(file_path: str) -> tuple[str, list[dict]]:
        """Parse document based on file extension"""
        ext = Path(file_path).suffix.lower()
        
        if ext == ".pdf":
            return ProductionDocumentParser.parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return ProductionDocumentParser.parse_docx(file_path)
        elif ext in [".png", ".jpg", ".jpeg"]:
            return ProductionDocumentParser.parse_image(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
